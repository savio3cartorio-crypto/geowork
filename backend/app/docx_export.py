"""Exportação do memorial descritivo em .docx (python-docx). Recebe as linhas já montadas pela interface."""

import io

from docx import Document
from docx.shared import Pt

ESTILOS = {
    # texto corrido: fonte de texto, parágrafo por linha
    "corrido": ("Times New Roman", 12),
    # tabular: fonte monoespaçada para manter as colunas alinhadas
    "tabular": ("Courier New", 9),
}


def gerar_memorial(titulo: str, linhas: list[str], formato: str) -> bytes:
    fonte, tamanho = ESTILOS[formato]
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = fonte
    estilo.font.size = Pt(tamanho)

    doc.add_heading(titulo, level=1)
    for linha in linhas:
        doc.add_paragraph(linha)

    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()
